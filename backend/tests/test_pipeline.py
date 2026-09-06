import asyncio
import base64
import copy
import io
import json
from pathlib import Path

import httpx
import pytest

from app.services.parsers import parse_artifact, sanitize_payload, sanitize_text
from app.services.pipeline import (
    PipelineError,
    compose_feedback,
    ingest_github_pr,
    load_github_example,
    run_evaluation,
    run_review_pipeline,
    validate_model_endpoint,
    validate_pr_url,
)


@pytest.fixture
def setup_review():
    assignment = {"id": "task1", "title": "HTTP server", "task_text": "Implement /ping"}
    rubric = {"version": 1, "criteria": [{"id": "c1", "title": "HTTP endpoint", "description": "Implement ping", "max_score": 2, "verification_mode": "llm"}]}
    config = {"id": "config-v1", "tasks": {"criterion_evaluation": {"model_id": "local", "prompt_template": "Оцени критерий", "params": {"temperature": 0.1, "max_retries": 0}}}, "thresholds": {"abstain": 0.6, "critic_confidence": 0.7}}
    models = [{"id": "local", "provider": "lm_studio", "base_url": "http://localhost:1234/v1", "model_name": "test-local", "enabled": True, "default_params": {"max_output_tokens": 1200}, "capabilities": {"json_mode": True}}]
    artifact = parse_artifact("main.go", b'package main\nfunc ping() { println("pong") }\n')
    artifact["public"] = True
    return assignment, rubric, config, models, [artifact]


def model_response(request, *, score=2, confidence=0.9):
    payload = json.loads(request.content)
    context = json.loads(payload["messages"][1]["content"])
    segment = context["segments"][0]
    raw = {"criterion_id": context["criterion"]["id"], "suggested_score": score, "confidence": confidence, "abstained": False, "reason": "Функция присутствует в исходнике.", "evidence": [{"segment_id": segment["id"], "quote": segment["text"].splitlines()[0]}], "annotations": []}
    return raw


def response(raw):
    return httpx.Response(200, json={"model": "test-local", "choices": [{"finish_reason": "stop", "message": {"content": json.dumps(raw)}}], "usage": {"prompt_tokens": 100, "completion_tokens": 50}})


def run_with_transport(setup_review, handler):
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await run_review_pipeline(*setup_review, client=client)
    return asyncio.run(invoke())


@pytest.mark.parametrize("url", ["http://github.com/org/repo/pull/1", "https://github.com.evil.test/org/repo/pull/1", "https://github.com/org/repo/pull/0", "https://github.com/org/repo/pull/1?url=http://localhost", "https://user@github.com/org/repo/pull/1", "https://github.com/org/repo/issues/1", "https://github.com/org/repo/pull/1/files", "https://github.com/org/repo/pull/1#fragment"])
def test_pr_url_rejects_noncanonical_or_unsafe_urls(url):
    with pytest.raises(PipelineError):
        validate_pr_url(url)


def test_pr_allowlist_is_enforced(monkeypatch):
    monkeypatch.setenv("APP_ENV", "dev_demo")
    assert validate_pr_url("https://github.com/Org/repo/pull/42", ["org/*"]) == ("Org/repo", 42)
    with pytest.raises(PipelineError, match="not_allowed"):
        validate_pr_url("https://github.com/other/repo/pull/42", ["org/*"])


def test_ingest_uses_immutable_head_and_does_not_follow_download_urls(monkeypatch):
    monkeypatch.setenv("APP_ENV", "dev_demo")
    monkeypatch.delenv("GITHUB_ALLOWED_REPOSITORIES", raising=False)
    head, blob = "a" * 40, "b" * 40
    urls = []

    def handler(request):
        urls.append(str(request.url))
        assert request.url.host == "api.github.com"
        path = request.url.path
        if path.endswith("/pulls/42"):
            return httpx.Response(200, json={"head": {"sha": head}, "base": {"sha": "c" * 40, "repo": {"private": False}}})
        if "/git/trees/" in path:
            assert head in path
            return httpx.Response(200, json={"sha": head, "truncated": False, "tree": [{"type": "blob", "mode": "100644", "path": "main.go", "sha": blob, "size": 13, "download_url": "http://169.254.169.254/secrets"}]})
        if "/git/blobs/" in path:
            return httpx.Response(200, json={"encoding": "base64", "content": base64.b64encode(b"package main\n").decode()})
        if path.endswith("/files"):
            return httpx.Response(200, json=[{"filename":"main.go", "status":"added", "sha":blob, "additions":1, "deletions":0, "patch":"@@ -0,0 +1 @@\n+package main"}])
        if path.endswith("/commits"):
            return httpx.Response(200, json=[])
        raise AssertionError(path)

    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await ingest_github_pr("https://github.com/org/repo/pull/42", client=client)
    result = asyncio.run(invoke())
    assert result["head_sha"] == head
    assert result["artifacts"][0]["segments"][0]["anchor"]["start"] == "line:1"
    assert result["artifacts"][0]["public"] is True
    assert len(urls) == 7


def test_ingest_rejects_truncated_tree(monkeypatch):
    monkeypatch.setenv("APP_ENV", "dev_demo")
    monkeypatch.delenv("GITHUB_ALLOWED_REPOSITORIES", raising=False)
    def handler(request):
        if request.url.path.endswith('/files'):
            return httpx.Response(200, json=[{'filename':'main.go','status':'added','sha':'b'*40}])
        if "/pulls/" in request.url.path:
            return httpx.Response(200, json={"head": {"sha": "a" * 40}})
        return httpx.Response(200, json={"truncated": True, "tree": []})
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            await ingest_github_pr("https://github.com/org/repo/pull/1", client=client)
    with pytest.raises(PipelineError, match="tree_truncated"):
        asyncio.run(invoke())


def test_source_content_preserves_line_numbers_across_chunks():
    source = "\n".join(f"line {i}" for i in range(1, 145))
    artifact = parse_artifact("source.py", source.encode())
    assert artifact["content"] == source
    assert artifact["segments"][1]["anchor"]["start"] == "line:61"
    assert artifact["segments"][1]["text"].startswith("line 61\n")


def test_html_strips_script_and_style_but_retains_visible_text():
    artifact = parse_artifact("answer.html", b"<style>secret CSS</style><p>Answer</p><script>fetch('evil')</script><div>Visible</div>")
    assert "Answer" in artifact["content"] and "Visible" in artifact["content"]
    assert "evil" not in artifact["content"] and "CSS" not in artifact["content"]


def test_ipynb_extracts_saved_outputs_without_execution():
    notebook = {"cells": [{"cell_type": "code", "source": ["raise RuntimeError('must not execute')"], "outputs": [{"text": ["saved output"], "data": {"image/png": "ignored"}}]}]}
    artifact = parse_artifact("answer.ipynb", json.dumps(notebook).encode())
    assert artifact["parse_status"] == "parsed"
    assert artifact["segments"][0]["anchor"]["start"] == "cell:1"
    assert artifact["segments"][1]["text"] == "saved output"
    assert "ignored" not in artifact["content"]


def test_docx_paragraph_and_table_anchors():
    from docx import Document
    document = Document()
    document.add_paragraph("Paragraph evidence")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
    buffer = io.BytesIO()
    document.save(buffer)
    artifact = parse_artifact("answer.docx", buffer.getvalue())
    assert artifact["parse_status"] == "parsed"
    assert artifact["segments"][0]["anchor"]["start"] == "paragraph:1"
    assert artifact["segments"][1]["anchor"]["start"] == "table:1:row:1:cell:1"


def test_xlsx_retains_formula_without_computing_it():
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active["A1"] = 7
    workbook.active["B1"] = "=A1*2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    artifact = parse_artifact("answer.xlsx", buffer.getvalue())
    assert artifact["parse_status"] == "parsed"
    assert artifact["segments"][1]["text"] == "=A1*2"
    assert artifact["segments"][1]["metadata"]["formula"] is True


def test_pdf_without_text_abstains():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    artifact = parse_artifact("scan.pdf", buffer.getvalue())
    assert artifact["parse_status"] == "needs_human"
    assert artifact["warning"] == "no_extractable_text"


@pytest.mark.parametrize("path,data", [("../outside.txt", b"unsafe"), ("unknown.bin", b"binary\x00data"), ("broken.docx", b"not a zip")])
def test_unsupported_artifacts_require_human(path, data):
    assert parse_artifact(path, data)["parse_status"] == "needs_human"


def test_sanitization_removes_tokens_email_and_phone():
    original = "OPENAI_API_KEY=sk-proj-abcdefghijklmnopqrstuv\ncontact=user@example.org\nphone=+7 (999) 123-45-67\npassword=unusual-secret\n"
    safe, counts = sanitize_text(original)
    assert "abcdefghijklmnopqrstuv" not in safe
    assert "user@example.org" not in safe
    assert "999" not in safe
    assert "unusual-secret" not in safe
    assert counts["secrets_found"] >= 2 and counts["pii_found"] >= 2


def test_nested_context_and_quoted_secrets_are_sanitized():
    safe = sanitize_payload({"reference": {"password": "contains spaces", "full_name": "Jane Doe"}, "text": '\"api_key\": \"a secret with spaces\"'})
    serialized = json.dumps(safe)
    assert "Jane Doe" not in serialized and "a secret with spaces" not in serialized and "contains spaces" not in serialized


def test_model_policy_prevents_sensitive_public_calls(monkeypatch):
    monkeypatch.setenv("ALLOW_PUBLIC_LLM", "true")
    monkeypatch.setenv("TEST_KEY", "test-token")
    model = {"base_url": "https://api.openai.com/v1", "provider": "openai", "model_name": "configured-model", "enabled": True, "api_key_env": "TEST_KEY"}
    with pytest.raises(PipelineError, match="sensitive_endpoint_blocked"):
        validate_model_endpoint(model, app_env="production")
    with pytest.raises(PipelineError, match="private_data"):
        validate_model_endpoint(model, app_env="dev_demo", public_data=False)
    assert validate_model_endpoint(model, app_env="dev_demo", public_data=True)["Authorization"] == "Bearer test-token"


def test_no_config_yields_no_fake_scores(setup_review):
    setup_review[2]["tasks"] = {}
    result = run_with_transport(setup_review, lambda request: pytest.fail("Must not call an unconfigured model"))
    assert result["status"] == "needs_human"
    assert result["error"].startswith("configuration_error:")
    assert result["criteria"][0]["suggested_score"] is None
    assert result["draft_total"] is None
    assert result["model_calls"] == []


def test_valid_review_is_anchored_and_records_actual_params(setup_review):
    received = []
    def handler(request):
        received.append(json.loads(request.content))
        return response(model_response(request))
    result = run_with_transport(setup_review, handler)
    assert result["status"] == "draft_ready"
    assert result["draft_total"] == 2
    assert result["criteria"][0]["evidence"][0]["valid"] is True
    assert "integrity" not in result
    call = result["model_calls"][0]
    assert call["model_id"] == "local" and call["input_tokens"] == 100
    assert call["effective_params"]["temperature"] == 0.1
    assert call["effective_params"]["max_tokens"] == 1200
    assert "prompt" not in call and "content" not in call
    assert len(received) == 1


@pytest.mark.parametrize("corruption", ["anchor", "quote", "score", "schema", "no_evidence"])
def test_invalid_evidence_or_score_never_produces_grade(setup_review, corruption):
    def handler(request):
        raw = model_response(request)
        if corruption == "anchor": raw["evidence"][0]["segment_id"] = "invented"
        if corruption == "quote": raw["evidence"][0]["quote"] = "nonexistent snippet"
        if corruption == "score": raw["suggested_score"] = 3
        if corruption == "schema": raw["unapproved_field"] = "value"
        if corruption == "no_evidence": raw["evidence"] = []
        return response(raw)
    result = run_with_transport(setup_review, handler)
    assert result["criteria"][0]["abstained"] is True
    assert result["draft_total"] is None


def test_critic_runs_only_for_low_confidence(setup_review):
    setup_review[2]["tasks"]["review_critic"] = {"model_id": "local", "prompt_template": "CRITIC", "params": {"max_retries": 0}}
    count = 0
    def handler(request):
        nonlocal count
        count += 1
        return response(model_response(request, confidence=0.4 if count == 1 else 0.95))
    result = run_with_transport(setup_review, handler)
    assert count == 2 and result["draft_total"] == 2
    assert [call["task_type"] for call in result["model_calls"]] == ["criterion_evaluation", "review_critic"]


def test_all_context_is_sanitized_before_model(setup_review):
    setup_review[0]["task_text"] += " Contact teacher@example.org. password=staff-secret"
    artifact = parse_artifact("secrets.txt", b"safe first line\nOPENAI_API_KEY=sk-proj-abcdefghijklmnopqrstuv\nstudent@example.org\n")
    artifact["public"] = True
    setup_review[4][:] = [artifact]
    def handler(request):
        request_text = request.content.decode()
        for forbidden in ["staff-secret", "teacher@example.org", "abcdefghijklmnopqrstuv", "student@example.org"]:
            assert forbidden not in request_text
        return response(model_response(request))
    assert run_with_transport(setup_review, handler)["draft_total"] == 2


def test_feedback_excludes_unaccepted_or_private_facts():
    review = {"criteria": [{"title": "Accepted score", "final_score": 2, "max_score": 3, "confirmed": True, "note": "Confirmed note"}, {"title": "Unconfirmed score", "final_score": 3, "max_score": 3}], "annotations": [{"status": "accepted", "visible_to_student": True, "message": "Confirmed fact", "advice": "Approved advice"}, {"status": "pending", "visible_to_student": True, "message": "Hallucinated allegation"}, {"status": "accepted", "visible_to_student": False, "message": "Private reviewer note"}], "integrity": {"message": "Must never enter feedback"}}
    composed = compose_feedback(review)
    assert "Accepted score: 2 / 3" in composed
    assert "Confirmed fact Approved advice" in composed
    assert all(text not in composed for text in ["Unconfirmed score", "Hallucinated", "Private", "Must never"])


def test_evaluation_limits_legacy_repeat_request_to_first_example(setup_review):
    assignment, rubric, config, models, artifacts = setup_review
    examples = []
    for level in ["weak", "medium", "good"]:
        artifact = parse_artifact("answer.txt", f"{level}\nsource evidence".encode())
        artifact["public"] = True
        examples.append({"level": level, "artifacts": [artifact]})
    def handler(request):
        context = json.loads(json.loads(request.content)["messages"][1]["content"])
        level = context["segments"][0]["text"].splitlines()[0]
        return response(model_response(request, score={"weak": 0, "medium": 1, "good": 2}[level]))
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await run_evaluation(assignment, rubric, config, models, examples, repetitions=3, client=client)
    result = asyncio.run(invoke())
    assert result["status"] == "completed" and len(result["outputs"]) == 1
    assert result["metrics"]["ordering_accuracy"] is None
    assert result["metrics"]["stability"]["good"]["stddev"] is None
    assert result["metrics"]["anchor_validity"] == 1
    assert result["metrics"]["input_tokens"] == 100
    assert result["outputs"][0]["raw_outputs"][0]["output"]["suggested_score"] == 0


def test_evaluation_does_not_invent_metrics_without_model(setup_review):
    assignment, rubric, config, models, artifacts = setup_review
    config["tasks"] = {}
    examples = [{"level": level, "artifacts": copy.deepcopy(artifacts)} for level in ["weak", "medium", "good"]]
    result = asyncio.run(run_evaluation(assignment, rubric, config, models, examples, repetitions=3))
    assert result["status"] == "failed"
    assert result["metrics"]["ordering_accuracy"] is None
    assert result["metrics"]["anchor_validity"] is None
    assert result["metrics"]["input_tokens"] is None
    assert result["metrics"]["abstain_rate"] == 1


def test_bundled_calibration_matches_git_blobs_without_network():
    fixture = json.loads((Path(__file__).parents[1] / "fixtures" / "go-task-1.json").read_text())
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: pytest.fail("Bundled references must not use network"))) as client:
            return [await load_github_example(example, client=client) for example in fixture["references"]]
    examples = asyncio.run(invoke())
    assert [len(example["artifacts"]) for example in examples] == [15, 7, 19]
    assert all(example["source"] == "bundled_verified_snapshot" for example in examples)


def test_reference_branch_is_pinned_before_loading():
    fixture = json.loads((Path(__file__).parents[1] / "fixtures" / "go-task-1.json").read_text())
    reference = {**fixture["references"][0], "ref": "main"}
    requests = []
    def handler(request):
        requests.append(request.url.path)
        assert request.url.path.endswith("/branches/main")
        return httpx.Response(200, json={"commit": {"sha": fixture["references"][0]["ref"]}})
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await load_github_example(reference, client=client)
    result = asyncio.run(invoke())
    assert len(requests) == 1
    assert result["head_sha"] == fixture["references"][0]["ref"]
    assert result["requested_ref"] == "main"


def test_optional_reference_is_context_not_student_evidence(setup_review):
    fixture = json.loads((Path(__file__).parents[1] / "fixtures" / "go-task-1.json").read_text())
    setup_review[0]["references"] = [{**fixture["references"][0], "type": "reference", "name": "Reference implementation"}]
    def handler(request):
        context = json.loads(json.loads(request.content)["messages"][1]["content"])
        assert context["reference_context"][0]["segments"]
        assert context["reference_context"][0]["head_sha"] == fixture["references"][0]["ref"]
        raw = model_response(request)
        reference_segment = context["reference_context"][0]["segments"][0]
        raw["evidence"] = [{"segment_id": reference_segment["id"], "quote": reference_segment["text"].splitlines()[0]}]
        return response(raw)
    result = run_with_transport(setup_review, handler)
    assert result["draft_total"] is None
    assert result["error"] == "invalid_source_anchor"


@pytest.fixture
def setup_gemma(setup_review, monkeypatch):
    monkeypatch.setenv("APP_ENV", "dev_demo")
    monkeypatch.setenv("ALLOW_PUBLIC_LLM", "true")
    monkeypatch.setenv("GEMINI_API_KEY", "test-gemini-secret-only-header")
    setup_review[3][0].update(provider="gemini", base_url="https://generativelanguage.googleapis.com/v1beta", model_name="gemma-4-31b-it", api_key_env="GEMINI_API_KEY", capabilities={"json_schema": False, "json_mode": False})
    setup_review[2]["tasks"]["criterion_evaluation"]["params"]["top_p"] = 0.85
    return setup_review


def gemma_output(request):
    payload = json.loads(request.content)
    context = json.loads(payload["contents"][0]["parts"][0]["text"])
    segment = context["segments"][0]
    return {"criterion_id": context["criterion"]["id"], "suggested_score": 2, "confidence": 0.9, "abstained": False, "reason": "Функция присутствует в исходнике.", "evidence": [{"segment_id": segment["id"], "quote": segment["text"].splitlines()[0]}], "annotations": []}


def gemma_response(raw, *, finish_reason="STOP", parts=None):
    return httpx.Response(200, json={"modelVersion": "gemma-4-31b-it-served", "candidates": [{"finishReason": finish_reason, "content": {"role": "model", "parts": parts if parts is not None else [{"text": json.dumps(raw)}]}}], "usageMetadata": {"promptTokenCount": 100, "candidatesTokenCount": 50, "thoughtsTokenCount": 7, "totalTokenCount": 157}})


def test_gemma_native_request_masks_context_and_discards_thoughts(setup_gemma):
    setup_gemma[0]["task_text"] += " Contact student@example.org; password=student-secret"
    received = []
    def handler(request):
        assert str(request.url) == "https://generativelanguage.googleapis.com/v1beta/models/gemma-4-31b-it:generateContent"
        assert request.headers["x-goog-api-key"] == "test-gemini-secret-only-header"
        assert "Authorization" not in request.headers
        assert not request.url.query
        payload = json.loads(request.content)
        received.append(payload)
        assert set(payload) == {"systemInstruction", "contents", "generationConfig"}
        assert payload["contents"][0]["role"] == "user"
        assert "Required JSON schema:" in payload["systemInstruction"]["parts"][0]["text"]
        assert payload["generationConfig"] == {"temperature": 0.1, "topP": 0.85, "maxOutputTokens": 1200}
        for secret in ["student@example.org", "student-secret", "test-gemini-secret-only-header"]:
            assert secret not in request.content.decode()
        raw = json.dumps(gemma_output(request))
        cut = len(raw) // 2
        return gemma_response(None, parts=[{"thought": True, "text": "private thinking, not JSON"}, {"text": raw[:cut]}, {"text": raw[cut:]}])
    result = run_with_transport(setup_gemma, handler)
    assert len(received) == 1 and result["draft_total"] == 2
    assert "private thinking" not in json.dumps(result)
    assert "test-gemini-secret-only-header" not in json.dumps(result)
    call = result["model_calls"][0]
    assert call["actual_model"] == "gemma-4-31b-it-served"
    assert call["finish_reason"] == "STOP"
    assert call["input_tokens"] == 100
    assert call["candidate_tokens"] == 50 and call["reasoning_tokens"] == 7
    assert call["output_tokens"] == 57 and call["total_tokens"] == 157
    assert call["effective_params"]["maxOutputTokens"] == 1200


@pytest.mark.parametrize("capability", ["json_schema", "json_mode"])
def test_gemini_structured_mode_requires_explicit_capability(setup_gemma, capability):
    setup_gemma[3][0]["capabilities"][capability] = True
    setup_gemma[3][0]["model_name"] = "models/gemma-4-31b-it"
    def handler(request):
        assert "/models/models/" not in request.url.path
        payload = json.loads(request.content)
        generation = payload["generationConfig"]
        assert generation["responseMimeType"] == "application/json"
        assert "response_format" not in payload and "responseSchema" not in generation
        if capability == "json_schema":
            schema = generation["responseJsonSchema"]
            assert schema["type"] == "object"
            assert "evidence" in schema["properties"]
            assert "Evidence" in schema["$defs"]
            assert "minLength" not in json.dumps(schema) and "default" not in json.dumps(schema)
        else:
            assert "responseJsonSchema" not in generation
        return gemma_response(gemma_output(request))
    assert run_with_transport(setup_gemma, handler)["draft_total"] == 2


def test_gemma_does_not_assume_json_mode_when_capability_is_absent(setup_gemma):
    setup_gemma[3][0]["capabilities"] = {}
    def handler(request):
        generation = json.loads(request.content)["generationConfig"]
        assert "responseMimeType" not in generation and "responseJsonSchema" not in generation
        return gemma_response(gemma_output(request))
    assert run_with_transport(setup_gemma, handler)["draft_total"] == 2


def test_gemma_token_override_uses_native_parameter(setup_gemma):
    task_params = setup_gemma[2]["tasks"]["criterion_evaluation"]["params"]
    task_params.update(max_tokens=900, stop="END", reasoning_effort="minimal")
    def handler(request):
        generation = json.loads(request.content)["generationConfig"]
        assert generation["maxOutputTokens"] == 900
        assert generation["stopSequences"] == ["END"]
        assert generation["thinkingConfig"] == {"thinkingLevel": "minimal"}
        assert "max_tokens" not in generation and "reasoning_effort" not in generation
        return gemma_response(gemma_output(request))
    assert run_with_transport(setup_gemma, handler)["draft_total"] == 2


def test_gemma_missing_key_is_configuration_error_without_request(setup_gemma, monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY")
    result = run_with_transport(setup_gemma, lambda request: pytest.fail("Must not call Gemini without credentials"))
    assert result["error"] == "configuration_error:credential_missing"
    assert result["draft_total"] is None
    assert result["model_calls"] == []


def test_gemini_preserves_sensitive_data_policy(setup_gemma):
    setup_gemma[4][0]["public"] = False
    result = run_with_transport(setup_gemma, lambda request: pytest.fail("Private artifacts must not reach public Gemini"))
    assert result["error"] == "configuration_error:private_data_public_endpoint_blocked"
    with pytest.raises(PipelineError, match="sensitive_endpoint_blocked"):
        validate_model_endpoint(setup_gemma[3][0], app_env="production")


@pytest.mark.parametrize("finish_reason,error", [("MAX_TOKENS", "model_incomplete_output"), (None, "model_incomplete_output"), ("OTHER", "model_incomplete_output"), ("SAFETY", "model_response_blocked"), ("RECITATION", "model_response_blocked"), ("SPII", "model_response_blocked"), ("PROHIBITED_CONTENT", "model_response_blocked")])
def test_gemma_rejects_non_successful_finish_even_with_valid_json(setup_gemma, finish_reason, error):
    result = run_with_transport(setup_gemma, lambda request: gemma_response(gemma_output(request), finish_reason=finish_reason))
    assert result["error"] == error
    assert result["draft_total"] is None
    assert result["model_calls"][0]["total_tokens"] == 157


def test_gemma_prompt_block_does_not_retry_or_call_critic(setup_gemma):
    setup_gemma[2]["tasks"]["criterion_evaluation"]["params"]["max_retries"] = 2
    setup_gemma[2]["tasks"]["review_critic"] = {"model_id": "local", "prompt_template": "Check the prior draft", "params": {"max_retries": 2}}
    seen = []
    def handler(request):
        seen.append(request)
        return httpx.Response(200, json={"promptFeedback": {"blockReason": "SAFETY"}, "usageMetadata": {"promptTokenCount": 99}})
    result = run_with_transport(setup_gemma, handler)
    assert result["error"] == "model_response_blocked"
    assert len(seen) == 1 and len(result["model_calls"]) == 1
    assert result["model_calls"][0]["input_tokens"] == 99


@pytest.mark.parametrize("parts,error", [([{"thought": True, "text": "Only hidden reasoning"}], "model_empty_output"), ([{"functionCall": {"name": "run_code"}}], "model_unexpected_tool_output"), (["malformed part"], "model_invalid_response")])
def test_gemma_rejects_missing_text_or_tool_results(setup_gemma, parts, error):
    result = run_with_transport(setup_gemma, lambda request: gemma_response(None, parts=parts))
    assert result["error"] == error and result["draft_total"] is None


def test_gemma_invalid_json_retries_within_configured_limit(setup_gemma, monkeypatch):
    setup_gemma[2]["tasks"]["criterion_evaluation"]["params"]["max_retries"] = 1
    async def skip_wait(_):
        return None
    monkeypatch.setattr("app.services.pipeline.asyncio.sleep", skip_wait)
    seen = []
    def handler(request):
        seen.append(request)
        if len(seen) == 1:
            return gemma_response(None, parts=[{"text": "invalid JSON"}])
        return gemma_response(gemma_output(request))
    result = run_with_transport(setup_gemma, handler)
    assert len(seen) == 2 and result["draft_total"] == 2
    assert [call["error_code"] for call in result["model_calls"]] == ["model_invalid_json", None]


def test_gemma_exhausted_json_retries_abstain(setup_gemma, monkeypatch):
    setup_gemma[2]["tasks"]["criterion_evaluation"]["params"]["max_retries"] = 2
    async def skip_wait(_):
        return None
    monkeypatch.setattr("app.services.pipeline.asyncio.sleep", skip_wait)
    result = run_with_transport(setup_gemma, lambda request: gemma_response(None, parts=[{"text": "invalid JSON"}]))
    assert len(result["model_calls"]) == 3
    assert result["error"] == "model_invalid_json"
    assert result["draft_total"] is None and result["criteria"][0]["abstained"] is True


def test_gemma_provider_error_body_is_not_recorded(setup_gemma):
    result = run_with_transport(setup_gemma, lambda request: httpx.Response(401, json={"error": {"message": "Rejected test-gemini-secret-only-header"}}))
    assert result["error"] == "upstream_http_401"
    assert "test-gemini-secret-only-header" not in json.dumps(result)


@pytest.mark.parametrize("surrounding", ["Accidentally pasted {} in a document", "https://generativelanguage.googleapis.com/v1beta/models?key={}&alt=json", '"{}"'])
def test_sanitizer_masks_google_key_in_plain_text_or_url(surrounding):
    key = "AIza" + "A" * 34 + "-"
    safe, counts = sanitize_text(surrounding.format(key))
    assert key not in safe
    assert "[REDACTED_SECRET]" in safe
    assert counts["secrets_found"] >= 1


@pytest.mark.parametrize('fallback', ['truncated', 'large_response', 'snapshot_limit', 'symlink', 'changed_head'])
def test_large_pr_fallback_is_pinned_bounded_and_skips_links(monkeypatch, fallback):
    monkeypatch.setenv('APP_ENV', 'dev_demo')
    monkeypatch.delenv('GITHUB_ALLOWED_REPOSITORIES', raising=False)
    head, folder, blob = 'a'*40, 'b'*40, 'c'*40
    pr = {'head': {'sha': head}, 'base': {'sha': 'd'*40, 'repo': {'private': False}}, 'changed_files': 2}
    calls = []
    metadata_calls = 0
    def handler(request):
        nonlocal metadata_calls
        calls.append(request)
        assert request.url.host == 'api.github.com'
        path = request.url.path
        if path.endswith('/pulls/42'):
            metadata_calls += 1
            return httpx.Response(200, json={**pr, 'head': {'sha': 'e'*40}} if fallback == 'changed_head' and metadata_calls > 1 else pr)
        if path.endswith('/files'):
            return httpx.Response(200, json=[{'filename': 'src/main.go', 'status': 'modified', 'sha': blob, 'additions':1, 'deletions':0, 'patch':'@@ -0,0 +1 @@\n+package main', 'download_url': 'https://evil.test/file'}, {'filename': 'old.go', 'status': 'removed', 'sha': 'f'*40}])
        if path.endswith('/commits'):
            return httpx.Response(200, json=[])
        if request.url.params.get('recursive'):
            if fallback == 'large_response':
                return httpx.Response(200, content=b' ' * 12_000_001)
            if fallback == 'snapshot_limit':
                return httpx.Response(200, json={'sha':head,'tree':[{'path':'src/main.go','type':'blob','mode':'100644','sha':blob,'size':13},{'path':'large.bin','type':'blob','mode':'100644','sha':'f'*40,'size':20_000_001}]})
            return httpx.Response(200, json={'sha':head,'truncated':True,'tree':[]})
        if path.endswith('/git/trees/'+head):
            return httpx.Response(200, json={'sha':head,'tree':[{'path':'src','type':'tree','mode':'040000','sha':folder}]})
        if path.endswith('/git/trees/'+folder):
            return httpx.Response(200, json={'sha':folder,'tree':[{'path':'main.go','type':'blob','mode':'120000' if fallback=='symlink' else '100644','sha':blob,'size':13}]})
        if path.endswith('/git/blobs/'+blob):
            assert fallback != 'symlink'
            return httpx.Response(200, json={'sha':blob,'encoding':'base64','content':base64.b64encode(b'package main\n').decode()})
        raise AssertionError(path)
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await ingest_github_pr('https://github.com/org/repo/pull/42', client=client, max_files=2, max_bytes=10_000_000)
    if fallback == 'changed_head':
        with pytest.raises(PipelineError, match='github_pr_changed_during_ingest'):
            asyncio.run(invoke())
        assert not any('/git/blobs/' in str(r.url) for r in calls)
        return
    result = asyncio.run(invoke())
    assert result['snapshot_scope'] == 'pull_request' and result['snapshot_complete'] is False
    assert result['head_sha'] == head
    assert [a['path'] for a in result['artifacts']] == ['src/main.go', 'old.go']
    assert result['artifacts'][0]['snapshot_complete'] is False
    assert result['artifacts'][0]['parse_status'] == ('needs_human' if fallback == 'symlink' else 'parsed')
    assert any(item['reason']=='removed_from_head' for item in result['context_limitations']['omitted_files'])
    assert all('evil.test' not in str(r.url) for r in calls)


@pytest.mark.parametrize('score,sufficient', [(0, True), (2, False), (2, None)])
def test_partial_snapshot_never_scores_missing_context(setup_review, score, sufficient):
    setup_review[4][0]['snapshot_complete'] = False
    def handler(request):
        context = json.loads(request.content)['messages'][1]['content']
        assert json.loads(context)['snapshot_coverage']['complete'] is False
        raw = model_response(request, score=score)
        raw['context_sufficient'] = sufficient
        return response(raw)
    result = run_with_transport(setup_review, handler)
    assert result['criteria'][0]['suggested_score'] is None
    assert result['draft_total'] is None and result['status'] == 'needs_human'


def test_quota_exhaustion_stops_retries_critic_and_remaining_criteria(setup_gemma):
    assignment, rubric, config, models, artifacts = setup_gemma
    rubric['criteria'] = [{**rubric['criteria'][0], 'id': f'c{i}'} for i in range(3)]
    config['tasks']['criterion_evaluation']['params']['max_retries'] = 2
    config['tasks']['review_critic'] = copy.deepcopy(config['tasks']['criterion_evaluation'])
    result = run_with_transport(setup_gemma, lambda request: httpx.Response(429, json={'error': {'message': 'private-provider-detail'}}))
    assert len(result['model_calls']) == 1
    assert result['error'] == 'upstream_http_429'
    assert len(result['criteria']) == 3
    assert all(item['suggested_score'] is None for item in result['criteria'])
    assert 'private-provider-detail' not in json.dumps(result)


def test_valid_lm_studio_abstention_preserves_reason_and_ignores_reasoning(setup_review):
    setup_review[2]['tasks']['review_critic'] = {'model_id': 'local', 'params': {'max_retries': 0}}
    reason = 'Нет сведений об обработке SIGINT/SIGTERM и завершении через context.'
    calls = []
    def handler(request):
        calls.append(request)
        raw = {'criterion_id': 'c1', 'suggested_score': None, 'confidence': 0.0, 'abstained': True,
               'context_sufficient': False, 'reason': reason, 'evidence': [], 'annotations': []}
        return httpx.Response(200, json={'choices': [{'finish_reason': 'stop', 'message': {'content': json.dumps(raw), 'reasoning': 'PRIVATE REASONING', 'tool_calls': []}}]})
    result = run_with_transport(setup_review, handler)
    assert len(calls) == 1
    assert result['error'] is None
    assert result['criteria'][0]['reason'] == reason
    assert result['criteria'][0]['suggested_score'] is None
    assert 'PRIVATE REASONING' not in json.dumps(result)


def test_legacy_demo_source_anchors_work_with_real_evidence(setup_review):
    setup_review[4][:] = [{'id': 'demo-artifact', 'path': 'main.go', 'parse_status': 'parsed', 'public': True,
                          'segments': [{'id': 'line-1', 'path': 'main.go', 'anchor': 'line:1', 'text': 'package main'}]}]
    result = run_with_transport(setup_review, lambda request: response(model_response(request)))
    assert result['draft_total'] == 2
    evidence = result['criteria'][0]['evidence'][0]
    assert evidence['artifact_id'] == 'demo-artifact'
    assert evidence['anchor']['start'] == 'line:1'


def test_lm_studio_schema_validation_retries_without_leaking_response(setup_review, monkeypatch):
    setup_review[3][0]['capabilities'] = {'json_schema': True}
    setup_review[2]['tasks']['criterion_evaluation']['params']['max_retries'] = 1
    async def no_sleep(_): pass
    monkeypatch.setattr('app.services.pipeline.asyncio.sleep', no_sleep)
    calls = 0
    def handler(request):
        nonlocal calls
        calls += 1
        payload = json.loads(request.content)
        assert payload['response_format']['json_schema']['strict'] is True
        raw = model_response(request)
        if calls == 1: raw['confidence'] = 'SECRET-INVALID-VALUE'
        return response(raw)
    result = run_with_transport(setup_review, handler)
    assert result['draft_total'] == 2 and calls == 2
    assert result['model_calls'][0]['error_code'] == 'invalid_criterion_schema'
    assert result['model_calls'][0]['validation_errors'][0]['field'] == 'confidence'
    assert 'SECRET-INVALID-VALUE' not in json.dumps(result)


def test_channel_error_maps_to_actionable_code(setup_review):
    def handler(request):
        return httpx.Response(500, json={"error": {"message": "Channel Error", "type": "server_error"}})

    result = run_with_transport(setup_review, handler)
    assert result["status"] == "needs_human"
    assert result["error"] == "upstream_channel_error"
    assert result["model_calls"][0]["error_code"] == "upstream_channel_error"

@pytest.mark.parametrize('stop_reason,expected', [('end_turn', 'draft_ready'), ('max_tokens', 'needs_human'), ('refusal', 'needs_human')])
def test_anthropic_messages(setup_review, monkeypatch, stop_reason, expected):
    monkeypatch.setenv('ANTHROPIC_API_KEY', 'test-anthropic-key')
    monkeypatch.setenv('ALLOW_PUBLIC_LLM', 'true')
    setup_review[3][0].update(provider='anthropic', base_url='https://api.anthropic.com/v1', model_name='claude-sonnet-4-6', api_key_env='ANTHROPIC_API_KEY')
    def handler(request):
        payload = json.loads(request.content)
        assert str(request.url) == 'https://api.anthropic.com/v1/messages'
        assert request.headers['x-api-key'] == 'test-anthropic-key'
        assert request.headers['anthropic-version'] == '2023-06-01'
        assert 'authorization' not in request.headers
        assert payload['max_tokens'] == 1200 and 'response_format' not in payload
        assert payload['system'] and payload['messages'][0]['role'] == 'user'
        context = json.loads(payload['messages'][0]['content'])
        segment = context['segments'][0]
        raw = {'criterion_id':'c1', 'suggested_score':2, 'confidence':0.95, 'abstained':False, 'reason':'Есть реализация', 'evidence':[{'segment_id':segment['id'], 'quote':segment['text'].splitlines()[0]}], 'annotations':[]}
        return httpx.Response(200, json={'model':'claude-sonnet-4-6', 'stop_reason':stop_reason, 'content':[{'type':'thinking','thinking':'private reasoning'}, {'type':'text','text':json.dumps(raw)}], 'usage':{'input_tokens':100,'output_tokens':50}})
    result = run_with_transport(setup_review, handler)
    assert result['status'] == expected
    assert 'test-anthropic-key' not in json.dumps(result)
    assert 'private reasoning' not in json.dumps(result)
    assert result['model_calls'][0]['total_tokens'] == 150


def test_anthropic_shared_agent_sampling_params():
    from app.services.pipeline import _effective_params
    model = {'provider':'anthropic', 'default_params':{'temperature':0.2, 'max_output_tokens':4000}}
    task = {'params':{'temperature':0.2, 'top_p':0.91, 'max_output_tokens':4096}}
    original = copy.deepcopy(task)
    params, _, _ = _effective_params(model, task)
    assert params == {'temperature':0.2, 'max_tokens':4096}
    assert task == original
    other, _, _ = _effective_params({**model, 'provider':'lm_studio'}, task)
    assert other['top_p'] == 0.91
    only_top_p, _, _ = _effective_params({'provider':'anthropic'}, {'params':{'top_p':0.9}})
    assert only_top_p['top_p'] == 0.9 and 'temperature' not in only_top_p


def test_eval_progress_arrives_before_next_example(setup_review):
    assignment, rubric, config, models, artifacts = setup_review
    examples = [{'level': level, 'artifacts': artifacts} for level in ['weak', 'medium', 'good']]
    events = []
    def handler(request):
        assert events[-1]['event'] == 'stage_started'
        assert events[-1]['repetition'] == 1
        return response(model_response(request))
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await run_evaluation(assignment, rubric, config, models, examples, repetitions=1, client=client, progress=events.append)
    result = asyncio.run(invoke())
    completed = [e for e in events if e['event'] == 'example_completed']
    assert [e['completed'] for e in completed] == [1]
    assert all(e['total'] == 1 for e in completed)
    assert [e['output'] for e in completed] == result['outputs']
    assert len([e for e in events if e['event'] == 'stage_completed']) == 1
    assert not any(e.get('level') == 'medium' for e in events)


@pytest.mark.parametrize('levels', [['good'], ['weak','good'], ['good','good']])
def test_eval_accepts_partial_and_repeated_categories(setup_review, levels):
    assignment, rubric, config, models, artifacts = setup_review
    examples = [{'level':level,'artifacts':artifacts} for level in levels]
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request: response(model_response(request)))) as client:
            return await run_evaluation(assignment,rubric,config,models,examples,client=client)
    result=asyncio.run(invoke())
    assert result['status']=='completed'
    assert len(result['outputs'])==1
    assert result['metrics']['example_count']==1
    assert result['metrics']['ordering_accuracy'] is None
    assert result['metrics']['anchor_validity']==1
    assert [o['example_index'] for o in result['outputs']]==[0]


def test_eval_never_retries_failed_model_request(setup_review):
    assignment,rubric,config,models,artifacts=setup_review
    config['tasks']['criterion_evaluation']['params']['max_retries']=2
    requests=[]
    def handler(request):
        requests.append(request)
        return httpx.Response(503,json={'error':'unavailable'})
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await run_evaluation(assignment,rubric,config,models,[{'level':'good','artifacts':artifacts}]*3,repetitions=3,client=client)
    result=asyncio.run(invoke())
    assert len(requests)==1
    assert len(result['outputs'])==1
    assert config['tasks']['criterion_evaluation']['params']['max_retries']==2
