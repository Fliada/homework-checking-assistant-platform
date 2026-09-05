"""Bounded, non-executing artifact normalization with reproducible source anchors."""

from __future__ import annotations

import hashlib
import io
import json
import mimetypes
import re
import zipfile
from html.parser import HTMLParser
from pathlib import PurePosixPath
from typing import Any


MAX_FILE_BYTES = 5_000_000
MAX_EXTRACTED_CHARS = 200_000
MAX_SEGMENTS = 2_000
SKIP_DIRECTORIES = {".git", "node_modules", "vendor", ".venv", "venv", "__pycache__", "dist", "build", ".next", "mlruns"}
GENERATED_SUFFIXES = {".lock", ".sum", ".map", ".min.js", ".min.css"}


class ParseLimitError(ValueError):
    pass


def safe_path(path: str) -> bool:
    return bool(path) and len(path) <= 1024 and not path.startswith("/") and "\\" not in path and "\x00" not in path and not any(part in {"..", ".", ""} for part in path.split("/"))


def should_skip(path: str) -> bool:
    p = PurePosixPath(path)
    return bool(set(p.parts) & SKIP_DIRECTORIES) or any(path.endswith(s) for s in GENERATED_SUFFIXES)


def _check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        files = archive.infolist()
        if len(files) > 2_000 or sum(f.file_size for f in files) > 32_000_000:
            raise ParseLimitError("archive_expansion_limit")
        for entry in files:
            if entry.flag_bits & 1 or (entry.file_size > 1_000_000 and entry.file_size / max(1, entry.compress_size) > 200):
                raise ParseLimitError("unsafe_archive")


class _VisibleHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[str] = []
        self.hidden: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript", "template"}:
            self.hidden.append(tag)
        if tag in {"p", "div", "li", "h1", "h2", "h3", "br", "tr"} and not self.hidden:
            self.blocks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if self.hidden and self.hidden[-1] == tag:
            self.hidden.pop()
        if tag in {"p", "div", "li", "tr"} and not self.hidden:
            self.blocks.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.hidden:
            self.blocks.append(data)


def parse_artifact(path: str, data: bytes, *, artifact_id: str | None = None) -> dict[str, Any]:
    """Parse bytes without writing files, running code, resolving links or formulas."""
    digest = hashlib.sha256(data).hexdigest()
    artifact_id = artifact_id or hashlib.sha256((path + digest).encode()).hexdigest()[:24]
    result: dict[str, Any] = {"id": artifact_id, "path": path, "sha256": digest, "size": len(data), "media_type": mimetypes.guess_type(path)[0] or "application/octet-stream", "parse_status": "parsed", "segments": [], "content": "", "warnings": []}
    if not safe_path(path):
        result.update(parse_status="needs_human", warning="unsafe_path")
        return result
    if should_skip(path):
        result.update(parse_status="skipped", warning="generated_or_dependency_file")
        return result
    if len(data) > MAX_FILE_BYTES:
        result.update(parse_status="needs_human", warning="file_size_limit")
        return result
    total_chars = 0

    def add(text: str, kind: str, start: str, end: str | None = None, **metadata: Any) -> None:
        nonlocal total_chars
        if not text.strip():
            return
        total_chars += len(text)
        if total_chars > MAX_EXTRACTED_CHARS or len(result["segments"]) >= MAX_SEGMENTS:
            raise ParseLimitError("extracted_content_limit")
        anchor = {"artifact_id": artifact_id, "path": path, "anchor_type": kind, "start": start, "end": end or start, "quote": text[:500]}
        result["segments"].append({"id": f"{artifact_id}:{len(result['segments']) + 1}", "artifact_id": artifact_id, "path": path, "anchor_type": kind, "anchor": anchor, "text": text, "metadata": metadata})

    suffix = PurePosixPath(path).suffix.lower()
    try:
        if suffix == ".docx":
            from docx import Document
            _check_zip(data)
            document = Document(io.BytesIO(data))
            for i, paragraph in enumerate(document.paragraphs, 1):
                add(paragraph.text, "paragraph", f"paragraph:{i}")
            for table_index, table in enumerate(document.tables, 1):
                for row_index, row in enumerate(table.rows, 1):
                    for col_index, cell in enumerate(row.cells, 1):
                        add(cell.text, "cell", f"table:{table_index}:row:{row_index}:cell:{col_index}")
            result["warnings"].append("Изображения DOCX не анализируются.")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError("encrypted_pdf")
            if len(reader.pages) > 100:
                raise ParseLimitError("pdf_page_limit")
            for i, page in enumerate(reader.pages, 1):
                add(page.extract_text() or "", "page", f"page:{i}")
            result["warnings"].append("Извлечён только текстовый слой PDF; OCR отключён.")
        elif suffix == ".ipynb":
            notebook = json.loads(data.decode("utf-8-sig"))
            cells = notebook.get("cells", [])
            if not isinstance(cells, list) or len(cells) > MAX_SEGMENTS:
                raise ParseLimitError("notebook_cell_limit")
            for i, cell in enumerate(cells, 1):
                source = cell.get("source", "")
                add("".join(source) if isinstance(source, list) else str(source), "notebook_cell", f"cell:{i}", cell_type=cell.get("cell_type", "unknown"))
                for output_index, output in enumerate(cell.get("outputs", []), 1):
                    text = output.get("text") or output.get("data", {}).get("text/plain", "")
                    add("".join(text) if isinstance(text, list) else str(text), "notebook_cell", f"cell:{i}:output:{output_index}", cell_type="output")
            result["warnings"].append("Код notebook и сохранённые outputs не выполнялись; изображения не анализируются.")
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            _check_zip(data)
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in workbook.worksheets:
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > 100_000:
                        raise ParseLimitError("spreadsheet_cell_limit")
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value is not None:
                                add(str(cell.value), "cell", f"{sheet.title}!{cell.coordinate}", sheet=sheet.title, coordinate=cell.coordinate, formula=cell.data_type == "f")
            finally:
                workbook.close()
            result["warnings"].append("Значения и формулы извлечены без пересчёта; изображения и оформление не анализируются.")
        elif suffix in {".html", ".htm"}:
            parser = _VisibleHTML()
            parser.feed(data.decode("utf-8-sig"))
            for i, block in enumerate("".join(parser.blocks).splitlines(), 1):
                add(block.strip(), "paragraph", f"block:{i}")
        else:
            if b"\x00" in data:
                raise ValueError("binary_artifact")
            source = data.decode("utf-8-sig")
            result["content"] = source
            if source and sum(ord(c) < 32 and c not in "\n\r\t" for c in source) / len(source) > 0.01:
                raise ValueError("binary_artifact")
            lines = source.splitlines()
            # Modest overlapping chunks keep line anchors useful during criterion retrieval.
            for start in range(0, len(lines), 60):
                chunk = lines[start:start + 70]
                add("\n".join(chunk), "line", f"line:{start + 1}", f"line:{start + len(chunk)}")
            if result["media_type"] == "application/octet-stream":
                result["media_type"] = "text/plain"
        if not result["segments"]:
            result.update(parse_status="needs_human", warning="no_extractable_text")
    except ParseLimitError as exc:
        result.update(parse_status="needs_human", warning=str(exc))
    except (ValueError, TypeError, KeyError, UnicodeError, zipfile.BadZipFile):
        result.update(parse_status="needs_human", warning="unsupported_or_invalid_artifact")
    except ImportError:
        result.update(parse_status="needs_human", warning="parser_not_installed")
    except Exception:
        # Parser diagnostics may contain student content; persist only a stable code.
        result.update(parse_status="needs_human", warning="parser_error")
    if not result["content"]:
        result["content"] = "\n\n".join(segment["text"] for segment in result["segments"])
    return result


_SECRET_PATTERNS = [
    re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----[\s\S]*?-----END (?:RSA |EC |OPENSSH )?PRIVATE KEY-----"),
    re.compile(r"\b(?:sk-(?:proj-)?[A-Za-z0-9_-]{16,}|gh[pousr]_[A-Za-z0-9]{20,}|github_pat_[A-Za-z0-9_]{20,}|AKIA[A-Z0-9]{16})\b"),
    re.compile(r"(?<![A-Za-z0-9_-])AIza[A-Za-z0-9_-]{35}(?![A-Za-z0-9_-])"),
    re.compile(r'''(?im)(["']?[\w.-]*(?:password|passwd|secret|api[_-]?key|access[_-]?token|auth[_-]?token)[\w.-]*["']?\s*(?::=|=|:)\s*)(?:"[^"\n]*"|'[^'\n]*'|[^\s,;\n]+)'''),
    re.compile(r"(?i)\bBearer\s+[A-Za-z0-9._~+/-]{10,}={0,2}"),
    re.compile(r"(?i)\b(?:postgres(?:ql)?|mysql|mongodb(?:\+srv)?|redis)://[^\s@]+@"),
]
_PII_PATTERNS = [
    re.compile(r"\b[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+\b"),
    re.compile(r"(?<![\w])\+?\d[\d ()-]{8,}\d(?![\w])"),
    re.compile(r"(?i)(?:ФИО|full[ _]?name|passport|паспорт|СНИЛС|ИНН)\s*[=:]\s*[^\n,;]+"),
]


def sanitize_text(text: str) -> tuple[str, dict[str, int]]:
    counts = {"secrets_found": 0, "pii_found": 0}
    for pattern in _SECRET_PATTERNS:
        text, count = pattern.subn(lambda match: (match.group(1) if pattern.groups else "") + "[REDACTED_SECRET]", text)
        counts["secrets_found"] += count
    for pattern in _PII_PATTERNS:
        text, count = pattern.subn("[REDACTED_PII]", text)
        counts["pii_found"] += count
    return text, counts


def sanitize_payload(value: Any) -> Any:
    """Sanitize all text leaves, including expert context and model output."""
    if isinstance(value, str):
        return sanitize_text(value)[0]
    if isinstance(value, dict):
        sanitized = {}
        for key, item in value.items():
            if re.search(r"(?i)(?:password|passwd|secret|api[_-]?key|access[_-]?token|auth[_-]?token)", str(key)):
                sanitized[key] = "[REDACTED_SECRET]"
            elif re.fullmatch(r"(?i)(?:email|phone|telephone|full[_ ]?name|ФИО|паспорт|passport|СНИЛС|ИНН)", str(key)):
                sanitized[key] = "[REDACTED_PII]"
            else:
                sanitized[key] = sanitize_payload(item)
        return sanitized
    if isinstance(value, list):
        return [sanitize_payload(item) for item in value]
    return value
